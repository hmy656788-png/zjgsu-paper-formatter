import ast
import random
import subprocess
import sys
import tempfile
import unittest
import zlib
from dataclasses import replace
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from zipfile import ZIP_BZIP2, ZIP_DEFLATED, ZipFile

from docx import Document

import docx_validation
import format_paper


class SharedDocxValidationTestCase(unittest.TestCase):
    @staticmethod
    def build_document_bytes(text: str = "正文") -> bytes:
        stream = BytesIO()
        document = Document()
        document.add_paragraph(text)
        document.save(stream)
        return stream.getvalue()

    @classmethod
    def build_word_object_count_docx(
        cls,
        main_paragraphs: int,
        *,
        comment_paragraphs: int = 0,
        main_runs: int = 0,
        comment_runs: int = 0,
    ) -> bytes:
        """构造高压缩段落包，不经 python-docx 逐段物化。"""
        source = BytesIO(cls.build_document_bytes())
        output = BytesIO()
        namespace = docx_validation.WORDPROCESSINGML_NAMESPACE
        document_xml = (
            f'<w:document xmlns:w="{namespace}"><w:body>'
            + "<w:p/>" * main_paragraphs
            + ("<w:p>" + "<w:r/>" * main_runs + "</w:p>" if main_runs else "")
            + "</w:body></w:document>"
        ).encode()

        with ZipFile(source, "r") as archive, ZipFile(
            output,
            "w",
            compression=ZIP_DEFLATED,
        ) as target:
            for member in archive.infolist():
                payload = archive.read(member)
                if member.filename == "word/document.xml":
                    payload = document_xml
                elif (
                    member.filename == "[Content_Types].xml"
                    and (comment_paragraphs or comment_runs)
                ):
                    payload = payload.replace(
                        b"</Types>",
                        (
                            b'<Override PartName="/word/comments.xml" '
                            b'ContentType="application/vnd.openxmlformats-officedocument.'
                            b'wordprocessingml.comments+xml"/></Types>'
                        ),
                    )
                target.writestr(member, payload)

            if comment_paragraphs or comment_runs:
                comments_xml = (
                    f'<w:comments xmlns:w="{namespace}"><w:comment w:id="0">'
                    + "<w:p/>" * comment_paragraphs
                    + (
                        "<w:p>" + "<w:r/>" * comment_runs + "</w:p>"
                        if comment_runs
                        else ""
                    )
                    + "</w:comment></w:comments>"
                ).encode()
                target.writestr("word/comments.xml", comments_xml)

        return output.getvalue()

    @staticmethod
    def find_member_header_offsets(
        archive_bytes: bytes,
        member_name: str,
    ) -> tuple[int, int]:
        with ZipFile(BytesIO(archive_bytes)) as archive:
            member = archive.getinfo(member_name)
            central_offset = archive.start_dir

        while archive_bytes[central_offset : central_offset + 4] == b"PK\x01\x02":
            filename_length = int.from_bytes(
                archive_bytes[central_offset + 28 : central_offset + 30],
                "little",
            )
            extra_length = int.from_bytes(
                archive_bytes[central_offset + 30 : central_offset + 32],
                "little",
            )
            comment_length = int.from_bytes(
                archive_bytes[central_offset + 32 : central_offset + 34],
                "little",
            )
            filename_start = central_offset + 46
            filename_end = filename_start + filename_length
            flags = int.from_bytes(
                archive_bytes[central_offset + 8 : central_offset + 10],
                "little",
            )
            encoding = "utf-8" if flags & 0x800 else "cp437"
            if archive_bytes[filename_start:filename_end].decode(encoding) == member_name:
                return member.header_offset, central_offset
            central_offset = (
                filename_end + extra_length + comment_length
            )

        raise AssertionError(f"central directory entry not found: {member_name}")

    @classmethod
    def build_binary_member_docx(cls, payload: bytes) -> bytes:
        source = BytesIO(cls.build_document_bytes())
        output = BytesIO()
        with ZipFile(source, "r") as archive, ZipFile(
            output,
            "w",
            compression=ZIP_DEFLATED,
        ) as target:
            for member in archive.infolist():
                member_payload = archive.read(member)
                if member.filename == "[Content_Types].xml":
                    member_payload = member_payload.replace(
                        b"</Types>",
                        (
                            b'<Default Extension="bin" '
                            b'ContentType="application/octet-stream"/></Types>'
                        ),
                    )
                target.writestr(member.filename, member_payload)
            target.writestr("word/media/payload.bin", payload)
        return output.getvalue()

    @classmethod
    def build_ignored_deep_xml_docx(cls) -> bytes:
        source = BytesIO(cls.build_document_bytes())
        output = BytesIO()
        extra_name = "word/ignored-depth.xml"
        deep_xml = (
            b"<?xml version='1.0' encoding='UTF-8'?><root>"
            + b"<item>" * (docx_validation.MAX_DOCX_XML_DEPTH + 1)
            + b"</item>" * (docx_validation.MAX_DOCX_XML_DEPTH + 1)
            + b"</root>"
        )

        with ZipFile(source, "r") as archive, ZipFile(
            output,
            "w",
            compression=ZIP_DEFLATED,
        ) as target:
            for member in archive.infolist():
                payload = archive.read(member)
                if member.filename == "[Content_Types].xml":
                    payload = payload.replace(
                        b"</Types>",
                        (
                            b'<Override PartName="/word/ignored-depth.xml" '
                            b'ContentType="application/xml"/></Types>'
                        ),
                    )
                target.writestr(member, payload)
            target.writestr(extra_name, deep_xml)

        return output.getvalue()

    def test_shared_validator_module_has_only_stdlib_imports(self):
        module_path = Path(docx_validation.__file__)
        tree = ast.parse(module_path.read_text(encoding="utf-8"))
        imported_roots = set()
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                imported_roots.update(
                    alias.name.split(".", 1)[0] for alias in node.names
                )
            elif isinstance(node, ast.ImportFrom) and node.module:
                imported_roots.add(node.module.split(".", 1)[0])

        self.assertEqual(
            imported_roots,
            {
                "contextlib",
                "dataclasses",
                "io",
                "pathlib",
                "posixpath",
                "re",
                "stat",
                "urllib",
                "xml",
                "zipfile",
                "zlib",
            },
        )

    def test_stream_validator_rewinds_valid_and_invalid_streams(self):
        valid_stream = BytesIO(self.build_document_bytes())
        valid_stream.seek(9)
        self.assertTrue(docx_validation.is_valid_docx_stream(valid_stream))
        self.assertEqual(valid_stream.tell(), 0)

        invalid_stream = BytesIO(b"not a docx")
        invalid_stream.seek(3)
        self.assertFalse(docx_validation.is_valid_docx_stream(invalid_stream))
        self.assertEqual(invalid_stream.tell(), 0)

        invalid_limits_stream = BytesIO(b"not a docx")
        invalid_limits_stream.seek(4)
        self.assertFalse(
            docx_validation.is_valid_docx_stream(
                invalid_limits_stream,
                limits=None,
            )
        )
        self.assertEqual(invalid_limits_stream.tell(), 0)

    def test_orphan_relationship_part_is_ignored_like_python_docx(self):
        """Unreachable ``*.rels`` parts are ignored by the package loader."""
        source = BytesIO(self.build_document_bytes("孤立关系部件"))
        output = BytesIO()
        orphan_relationships = (
            b'<?xml version="1.0" encoding="UTF-8"?>'
            b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>'
        )
        with ZipFile(source, "r") as archive, ZipFile(
            output,
            "w",
            compression=ZIP_DEFLATED,
        ) as target:
            for member in archive.infolist():
                target.writestr(member, archive.read(member))
            target.writestr("word/_rels/orphan.xml.rels", orphan_relationships)

        package_bytes = output.getvalue()
        # python-docx traverses only relationships reachable from _rels/.rels,
        # so this valid-but-unreferenced part is ignored by the consumer.
        Document(BytesIO(package_bytes))
        self.assertTrue(
            docx_validation.is_valid_docx_stream(BytesIO(package_bytes))
        )

    def test_orphan_relationship_alias_checks_scale_linearly(self):
        source = BytesIO(self.build_document_bytes())
        output = BytesIO()
        orphan_count = 32
        empty_relationships = (
            b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>'
        )
        with ZipFile(source) as archive, ZipFile(
            output, "w", compression=ZIP_DEFLATED
        ) as target:
            for member in archive.infolist():
                target.writestr(member, archive.read(member))
            for index in range(orphan_count):
                target.writestr(
                    f"word/_rels/orphan{index}.xml.rels", empty_relationships
                )

        with ZipFile(BytesIO(output.getvalue())) as archive:
            member_names = set(archive.namelist())
            defaults, overrides = docx_validation.parse_docx_content_types(archive)
            with patch.object(
                docx_validation,
                "canonical_docx_part_name",
                wraps=docx_validation.canonical_docx_part_name,
            ) as canonicalize:
                self.assertEqual(
                    docx_validation.validate_docx_relationships(
                        archive, member_names, defaults, overrides
                    ),
                    "word/document.xml",
                )
            # Each package member is indexed once and each orphan is looked
            # up once; a nested scan of all members for every orphan regresses
            # this bound even on fast development machines.
            self.assertLessEqual(
                canonicalize.call_count, len(member_names) + orphan_count
            )

    def test_paragraph_budget_rejects_compressed_object_amplification(self):
        exact_input = BytesIO(
            self.build_word_object_count_docx(
                docx_validation.MAX_DOCX_TOTAL_PARAGRAPHS,
            )
        )
        excessive_input = BytesIO(
            self.build_word_object_count_docx(
                docx_validation.MAX_DOCX_TOTAL_PARAGRAPHS + 1,
            )
        )

        self.assertTrue(docx_validation.is_valid_docx_stream(exact_input))
        self.assertFalse(docx_validation.is_valid_docx_stream(excessive_input))
        self.assertEqual(exact_input.tell(), 0)
        self.assertEqual(excessive_input.tell(), 0)

        exact_generated = BytesIO(
            self.build_word_object_count_docx(
                docx_validation.MAX_GENERATED_DOCX_TOTAL_PARAGRAPHS,
            )
        )
        excessive_generated = BytesIO(
            self.build_word_object_count_docx(
                docx_validation.MAX_GENERATED_DOCX_TOTAL_PARAGRAPHS + 1,
            )
        )
        self.assertTrue(
            docx_validation.is_valid_docx_stream(
                exact_generated,
                limits=docx_validation.GENERATED_DOCX_LIMITS,
            )
        )
        self.assertFalse(
            docx_validation.is_valid_docx_stream(
                excessive_generated,
                limits=docx_validation.GENERATED_DOCX_LIMITS,
            )
        )

    def test_paragraph_budget_is_shared_across_word_story_parts(self):
        main_paragraphs = docx_validation.MAX_DOCX_TOTAL_PARAGRAPHS // 2
        exact_total = BytesIO(
            self.build_word_object_count_docx(
                main_paragraphs,
                comment_paragraphs=(
                    docx_validation.MAX_DOCX_TOTAL_PARAGRAPHS
                    - main_paragraphs
                ),
            )
        )
        excessive_total = BytesIO(
            self.build_word_object_count_docx(
                main_paragraphs,
                comment_paragraphs=(
                    docx_validation.MAX_DOCX_TOTAL_PARAGRAPHS
                    - main_paragraphs
                    + 1
                ),
            )
        )

        self.assertTrue(docx_validation.is_valid_docx_stream(exact_total))
        self.assertFalse(docx_validation.is_valid_docx_stream(excessive_total))

    def test_run_budget_rejects_compressed_object_amplification(self):
        exact_input = BytesIO(
            self.build_word_object_count_docx(
                0,
                main_runs=docx_validation.MAX_DOCX_TOTAL_RUNS,
            )
        )
        excessive_input = BytesIO(
            self.build_word_object_count_docx(
                0,
                main_runs=docx_validation.MAX_DOCX_TOTAL_RUNS + 1,
            )
        )
        self.assertTrue(docx_validation.is_valid_docx_stream(exact_input))
        self.assertFalse(docx_validation.is_valid_docx_stream(excessive_input))

        exact_generated = BytesIO(
            self.build_word_object_count_docx(
                0,
                main_runs=docx_validation.MAX_GENERATED_DOCX_TOTAL_RUNS,
            )
        )
        excessive_generated = BytesIO(
            self.build_word_object_count_docx(
                0,
                main_runs=docx_validation.MAX_GENERATED_DOCX_TOTAL_RUNS + 1,
            )
        )
        self.assertTrue(
            docx_validation.is_valid_docx_stream(
                exact_generated,
                limits=docx_validation.GENERATED_DOCX_LIMITS,
            )
        )
        self.assertFalse(
            docx_validation.is_valid_docx_stream(
                excessive_generated,
                limits=docx_validation.GENERATED_DOCX_LIMITS,
            )
        )

    def test_run_budget_is_shared_across_word_story_parts(self):
        main_runs = docx_validation.MAX_DOCX_TOTAL_RUNS // 2
        exact_total = BytesIO(
            self.build_word_object_count_docx(
                0,
                main_runs=main_runs,
                comment_runs=docx_validation.MAX_DOCX_TOTAL_RUNS - main_runs,
            )
        )
        excessive_total = BytesIO(
            self.build_word_object_count_docx(
                0,
                main_runs=main_runs,
                comment_runs=(
                    docx_validation.MAX_DOCX_TOTAL_RUNS - main_runs + 1
                ),
            )
        )
        self.assertTrue(docx_validation.is_valid_docx_stream(exact_total))
        self.assertFalse(docx_validation.is_valid_docx_stream(excessive_total))

    def test_direct_formatter_rejects_paragraph_amplification_before_loading(self):
        archive_bytes = self.build_word_object_count_docx(
            docx_validation.MAX_DOCX_TOTAL_PARAGRAPHS + 1,
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            input_path = Path(temp_dir) / "too-many-paragraphs.docx"
            output_path = Path(temp_dir) / "output.docx"
            input_path.write_bytes(archive_bytes)
            with patch.object(format_paper, "Document") as document_factory:
                result = format_paper.format_academic_paper(
                    str(input_path),
                    str(output_path),
                )

            self.assertFalse(result)
            document_factory.assert_not_called()
            self.assertFalse(output_path.exists())

    def test_direct_formatter_rejects_run_amplification_before_loading(self):
        archive_bytes = self.build_word_object_count_docx(
            0,
            main_runs=docx_validation.MAX_DOCX_TOTAL_RUNS + 1,
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            input_path = Path(temp_dir) / "too-many-runs.docx"
            output_path = Path(temp_dir) / "output.docx"
            input_path.write_bytes(archive_bytes)
            with patch.object(format_paper, "Document") as document_factory:
                result = format_paper.format_academic_paper(
                    str(input_path),
                    str(output_path),
                )

            self.assertFalse(result)
            document_factory.assert_not_called()
            self.assertFalse(output_path.exists())

    def test_raw_archive_limit_rejects_before_zip_parsing(self):
        stream = BytesIO(b"oversized")
        limits = replace(
            docx_validation.INPUT_DOCX_LIMITS,
            max_archive_bytes=4,
        )

        with patch.object(
            docx_validation.zipfile,
            "is_zipfile",
            side_effect=AssertionError("oversized archives must not be parsed"),
        ) as is_zipfile:
            self.assertFalse(
                docx_validation.is_valid_docx_stream(stream, limits=limits)
            )

        is_zipfile.assert_not_called()
        self.assertEqual(stream.tell(), 0)

    def test_malformed_custom_limits_fail_closed_without_parser_errors(self):
        payload = self.build_document_bytes()
        for field_name, value in (
            ("max_total_uncompressed_bytes", -1),
            ("max_non_xml_compression_ratio", 0),
            ("max_xml_depth", True),
        ):
            with self.subTest(field_name=field_name):
                limits = replace(
                    docx_validation.INPUT_DOCX_LIMITS,
                    **{field_name: value},
                )
                stream = BytesIO(payload)
                self.assertFalse(
                    docx_validation.is_valid_docx_stream(stream, limits=limits)
                )
                self.assertEqual(stream.tell(), 0)

    def test_valid_data_descriptor_and_forced_zip64_layouts_are_accepted(self):
        source_bytes = self.build_document_bytes("ZIP 布局兼容性")
        with ZipFile(BytesIO(source_bytes), "r") as source:
            payloads = [
                (member.filename, source.read(member))
                for member in source.infolist()
            ]

        class NonSeekableBuffer(BytesIO):
            def seekable(self):
                return False

            def seek(self, *_args, **_kwargs):
                raise OSError("non-seekable output")

        descriptor_output = NonSeekableBuffer()
        with ZipFile(
            descriptor_output,
            "w",
            compression=ZIP_DEFLATED,
        ) as target:
            for member_name, payload in payloads:
                target.writestr(member_name, payload)
        descriptor_bytes = descriptor_output.getvalue()
        with ZipFile(BytesIO(descriptor_bytes), "r") as archive:
            self.assertTrue(
                all(member.flag_bits & 0x08 for member in archive.infolist())
            )
        descriptor_stream = BytesIO(descriptor_bytes)
        descriptor_stream.seek(7)
        self.assertTrue(
            docx_validation.is_valid_docx_stream(descriptor_stream)
        )
        self.assertEqual(descriptor_stream.tell(), 0)

        zip64_output = BytesIO()
        with ZipFile(
            zip64_output,
            "w",
            compression=ZIP_DEFLATED,
            allowZip64=True,
        ) as target:
            for member_name, payload in payloads:
                with target.open(
                    member_name,
                    "w",
                    force_zip64=True,
                ) as target_member:
                    target_member.write(payload)
        zip64_bytes = zip64_output.getvalue()
        with ZipFile(BytesIO(zip64_bytes), "r") as archive:
            first_member = archive.infolist()[0]
            local_offset = first_member.header_offset
        self.assertEqual(
            zip64_bytes[local_offset + 18 : local_offset + 26],
            b"\xff" * 8,
        )
        self.assertTrue(
            docx_validation.is_valid_docx_stream(BytesIO(zip64_bytes))
        )

    def test_forged_bzip2_method_is_rejected_before_member_decompression(self):
        forged = bytearray(self.build_document_bytes())
        local_offset, central_offset = self.find_member_header_offsets(
            forged,
            "[Content_Types].xml",
        )
        forged[local_offset + 8 : local_offset + 10] = ZIP_BZIP2.to_bytes(
            2,
            "little",
        )
        forged[central_offset + 10 : central_offset + 12] = ZIP_BZIP2.to_bytes(
            2,
            "little",
        )
        with ZipFile(BytesIO(forged)) as archive:
            self.assertEqual(
                archive.getinfo("[Content_Types].xml").compress_type,
                ZIP_BZIP2,
            )

        stream = BytesIO(forged)
        stream.seek(11)
        with patch.object(
            docx_validation.zipfile.ZipFile,
            "open",
            side_effect=AssertionError(
                "unsupported compression reached member decompression"
            ),
        ) as archive_open:
            self.assertFalse(docx_validation.is_valid_docx_stream(stream))

        archive_open.assert_not_called()
        self.assertEqual(stream.tell(), 0)

    def test_unsupported_general_purpose_flags_are_rejected_before_decompression(self):
        forged = bytearray(self.build_binary_member_docx(b"safe binary payload"))
        target_name = "word/media/payload.bin"
        local_offset, central_offset = self.find_member_header_offsets(
            forged,
            target_name,
        )
        patched_data_flag = 0x20
        local_flags = int.from_bytes(
            forged[local_offset + 6 : local_offset + 8],
            "little",
        )
        central_flags = int.from_bytes(
            forged[central_offset + 8 : central_offset + 10],
            "little",
        )
        forged[local_offset + 6 : local_offset + 8] = (
            local_flags | patched_data_flag
        ).to_bytes(2, "little")
        forged[central_offset + 8 : central_offset + 10] = (
            central_flags | patched_data_flag
        ).to_bytes(2, "little")

        with ZipFile(BytesIO(forged)) as archive:
            target = archive.getinfo(target_name)
            self.assertEqual(target.flag_bits & patched_data_flag, patched_data_flag)
            with self.assertRaises(NotImplementedError):
                archive.read(target)

        stream = BytesIO(forged)
        with patch.object(
            docx_validation,
            "_verify_zip_member_output",
            side_effect=AssertionError(
                "unsupported flags reached member decompression"
            ),
        ) as verify_output:
            self.assertFalse(docx_validation.is_valid_docx_stream(stream))

        verify_output.assert_not_called()
        self.assertEqual(stream.tell(), 0)

    def test_forged_deflate_sizes_cannot_hide_output_beyond_budget(self):
        base_docx = self.build_document_bytes()
        with ZipFile(BytesIO(base_docx)) as archive:
            largest_base_member = max(
                member.file_size for member in archive.infolist()
            )
        configured_limit = largest_base_member + 1024
        payload = random.Random(0).randbytes(configured_limit + 1)
        forged = bytearray(self.build_binary_member_docx(payload))
        target_name = "word/media/payload.bin"
        local_offset, central_offset = self.find_member_header_offsets(
            forged,
            target_name,
        )

        declared_crc = zlib.crc32(payload[:configured_limit])
        forged[local_offset + 14 : local_offset + 18] = declared_crc.to_bytes(
            4,
            "little",
        )
        forged[local_offset + 22 : local_offset + 26] = configured_limit.to_bytes(
            4,
            "little",
        )
        forged[central_offset + 16 : central_offset + 20] = declared_crc.to_bytes(
            4,
            "little",
        )
        forged[central_offset + 24 : central_offset + 28] = configured_limit.to_bytes(
            4,
            "little",
        )

        # ZipExtFile trusts the forged central size, returns only that prefix,
        # and considers its forged prefix CRC valid. The shared validator must
        # instead inspect the complete raw DEFLATE stream.
        with ZipFile(BytesIO(forged)) as archive:
            target = archive.getinfo(target_name)
            self.assertEqual(target.compress_type, ZIP_DEFLATED)
            self.assertEqual(target.file_size, configured_limit)
            self.assertEqual(len(archive.read(target)), configured_limit)
            self.assertIsNone(archive.testzip())

        limits = replace(
            docx_validation.INPUT_DOCX_LIMITS,
            max_member_uncompressed_bytes=configured_limit,
        )
        stream = BytesIO(forged)
        stream.seek(17)
        self.assertFalse(
            docx_validation.is_valid_docx_stream(stream, limits=limits)
        )
        self.assertEqual(stream.tell(), 0)

    def test_loader_validates_and_loads_the_same_open_handle(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            input_path = Path(temp_dir) / "input.docx"
            input_path.write_bytes(self.build_document_bytes("同句柄加载"))
            real_validator = format_paper.is_valid_docx_stream
            real_document = format_paper.Document
            observed = {}

            def observe_validation(stream, *, limits):
                self.assertFalse(stream.closed)
                self.assertIs(limits, format_paper.INPUT_DOCX_LIMITS)
                observed["stream"] = stream
                return real_validator(stream, limits=limits)

            def observe_load(stream):
                self.assertIs(stream, observed["stream"])
                self.assertFalse(stream.closed)
                return real_document(stream)

            with (
                patch.object(
                    format_paper,
                    "is_valid_docx_stream",
                    side_effect=observe_validation,
                ),
                patch.object(
                    format_paper,
                    "Document",
                    side_effect=observe_load,
                ),
            ):
                document = format_paper._load_validated_document(input_path)

            self.assertEqual(document.paragraphs[0].text, "同句柄加载")
            self.assertTrue(observed["stream"].closed)

    def test_merge_loads_generated_intermediate_with_expanded_profile(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            cover_path = temp_path / "cover.docx"
            body_path = temp_path / "body.docx"
            output_path = temp_path / "merged.docx"
            for path, text in (
                (cover_path, "封面"),
                (body_path, "正文"),
            ):
                document = Document()
                document.add_paragraph(text)
                document.save(path)

            real_loader = format_paper._load_validated_document
            observed_limits = []

            def observe_load(
                path,
                *,
                limits=format_paper.INPUT_DOCX_LIMITS,
            ):
                observed_limits.append(limits)
                return real_loader(path, limits=limits)

            with patch.object(
                format_paper,
                "_load_validated_document",
                side_effect=observe_load,
            ):
                result = format_paper.merge_cover_and_body(
                    cover_path,
                    body_path,
                    output_path,
                )

            self.assertIsInstance(result, dict)
            self.assertEqual(
                observed_limits,
                [
                    format_paper.INPUT_DOCX_LIMITS,
                    format_paper.INPUT_DOCX_LIMITS,
                    format_paper.GENERATED_DOCX_LIMITS,
                ],
            )
            self.assertTrue(output_path.exists())

    def test_formatter_rejects_validator_invalid_ignored_xml_before_load(self):
        malformed = self.build_ignored_deep_xml_docx()
        with tempfile.TemporaryDirectory() as temp_dir:
            input_path = Path(temp_dir) / "deep-ignored.docx"
            output_path = Path(temp_dir) / "formatted.docx"
            input_path.write_bytes(malformed)

            # python-docx ignores this orphan part, proving validation—not the
            # parser—is what rejects the resource-exhaustion payload.
            self.assertEqual(len(Document(input_path).paragraphs), 1)
            with patch.object(
                format_paper,
                "Document",
                side_effect=AssertionError("invalid input reached python-docx"),
            ) as document_loader:
                result = format_paper.format_academic_paper(
                    str(input_path),
                    str(output_path),
                )

            self.assertFalse(result)
            document_loader.assert_not_called()
            self.assertFalse(output_path.exists())

    def test_invalid_merge_cover_prevents_body_formatting(self):
        malformed = self.build_ignored_deep_xml_docx()
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            cover_path = temp_path / "cover.docx"
            body_path = temp_path / "body.docx"
            output_path = temp_path / "merged.docx"
            cover_path.write_bytes(malformed)
            body_path.write_bytes(self.build_document_bytes("正文"))

            with patch.object(
                format_paper,
                "format_academic_paper",
                side_effect=AssertionError("body formatting must not start"),
            ) as body_formatter:
                result = format_paper.merge_cover_and_body(
                    str(cover_path),
                    str(body_path),
                    str(output_path),
                )

            self.assertFalse(result)
            body_formatter.assert_not_called()
            self.assertFalse(output_path.exists())

    def test_concat_rejects_each_invalid_input_with_its_label(self):
        malformed = self.build_ignored_deep_xml_docx()
        valid = self.build_document_bytes("有效文档")
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            first_path = temp_path / "first.docx"
            second_path = temp_path / "second.docx"
            output_path = temp_path / "merged.docx"

            for invalid_position, expected_label in (
                ("first", "第一个文档"),
                ("second", "第二个文档"),
            ):
                with self.subTest(invalid_position=invalid_position):
                    first_path.write_bytes(
                        malformed if invalid_position == "first" else valid
                    )
                    second_path.write_bytes(
                        malformed if invalid_position == "second" else valid
                    )
                    with self.assertRaises(
                        format_paper.DocumentConcatError
                    ) as caught:
                        format_paper.concatenate_documents(
                            str(first_path),
                            str(second_path),
                            str(output_path),
                        )

                    self.assertIn(expected_label, str(caught.exception))
                    self.assertFalse(output_path.exists())

    def test_cli_rejects_validator_invalid_input_without_output(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            input_path = temp_path / "deep-ignored.docx"
            output_path = temp_path / "formatted.docx"
            input_path.write_bytes(self.build_ignored_deep_xml_docx())
            script_path = Path(format_paper.__file__)

            completed = subprocess.run(
                [
                    sys.executable,
                    str(script_path),
                    str(input_path),
                    str(output_path),
                ],
                capture_output=True,
                text=True,
                timeout=20,
                check=False,
            )

            self.assertEqual(completed.returncode, 1)
            diagnostic = completed.stdout + completed.stderr
            self.assertIn("document failed DOCX safety validation", diagnostic)
            self.assertNotIn("Traceback", diagnostic)
            self.assertFalse(output_path.exists())


if __name__ == "__main__":
    unittest.main()
